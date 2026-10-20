#!/usr/bin/env python3
"""Bundle the slide PNGs and the two wiring drawings into a 16:9 .pptx.
Needs python-pptx. Usage: build_deck.py OUT.pptx SLIDE_PNG... [-- FIT_PNG...]"""
import sys
from pptx import Presentation
from pptx.util import Emu
from PIL import Image

args = sys.argv[2:]
full, fit = (args[:args.index("--")], args[args.index("--") + 1:]) if "--" in args else (args, [])
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333 × 7.5 in
blank = prs.slide_layouts[6]
for p in full:
    prs.slides.add_slide(blank).shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height)
for p in fit:                                    # letterbox non-16:9 drawings on a white page
    w, h = Image.open(p).size
    sw, sh = prs.slide_width, prs.slide_height
    scale = min(sw / w, sh / h)
    pw, ph = int(w * scale), int(h * scale)
    prs.slides.add_slide(blank).shapes.add_picture(p, (sw - pw) // 2, (sh - ph) // 2, pw, ph)
prs.save(sys.argv[1])
